import csv
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import store
from app.main import app


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "ROOT", tmp_path)
    with TestClient(app) as test_client:
        yield test_client


def upload_csv(client, text, filename="客户.csv"):
    return client.post("/api/datasets", files={"file": (filename, text.encode("utf-8-sig"), "text/csv")})


def test_full_import_preview_task_exports_and_immutable_source(client):
    identity = upload_csv(client, "姓名,城市\n 张三 ,\n张三,未知\n李四,济南\n").json()["id"]
    body = {"name": "客户清洗", "rules": [
        {"id": "1", "type": "trim"},
        {"id": "2", "type": "fill", "fields": ["城市"], "options": {"value": "未知"}},
        {"id": "3", "type": "dedupe"},
    ]}
    preview = client.post(f"/api/datasets/{identity}/preview", json=body)
    assert preview.status_code == 200
    assert preview.json()["after"]["rows"] == 2
    response = client.post(f"/api/datasets/{identity}/tasks", json=body)
    assert response.status_code == 202
    task_id = response.json()["id"]
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "success"
    assert task["summary"]["changes"] == preview.json()["total_changes"]
    result = client.get(f"/api/tasks/{task_id}/result").json()
    assert result["total"] == 2
    assert result["items"] == preview.json()["rows"]
    source = client.get(f"/api/datasets/{identity}").json()
    assert source["total"] == 3
    assert source["rows"][0]["姓名"] == " 张三 "
    csv_result = client.get(f"/api/tasks/{task_id}/export")
    assert list(csv.reader(io.StringIO(csv_result.content.decode("utf-8-sig")))) == [["姓名", "城市"], ["张三", "未知"], ["李四", "济南"]]
    excel_result = client.get(f"/api/tasks/{task_id}/export?format=xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(excel_result.content))
    assert list(workbook.active.values)[1] == ("张三", "未知")
    assert client.get(f"/api/tasks/{task_id}/export?kind=changes").status_code == 200
    assert client.get(f"/api/tasks/{task_id}/export?kind=issues").status_code == 200
    assert client.get(f"/api/tasks/{task_id}/export?kind=report").json()["质量统计"]["清洗后"]["rows"] == 2
    assert client.get(f"/api/datasets/{identity}?search=济南").json()["total"] == 1
    assert client.get(f"/api/tasks/{task_id}/result?view=changes&page_size=1").json()["total"] == 3


@pytest.mark.parametrize("text", ["姓名,姓名\n甲,乙", "姓名,\n甲,乙", "_row_id,姓名\n1,甲", "姓名,城市\n甲", "姓名,城市\n"])
def test_import_rejects_ambiguous_or_empty_tables(client, text):
    response = upload_csv(client, text)
    assert response.status_code == 400


def test_xlsx_and_gb18030_preserve_phone_leading_zeroes(client):
    workbook = openpyxl.Workbook()
    workbook.active.append(["姓名", "编号"])
    workbook.active.append(["测试", "00123"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = client.post("/api/datasets", files={"file": ("测试.xlsx", buffer.getvalue())})
    assert response.status_code == 201
    assert client.get(f"/api/datasets/{response.json()['id']}").json()["rows"][0]["编号"] == "00123"
    response = client.post("/api/datasets", files={"file": ("测试.csv", "姓名,编号\n测试,00123".encode("gb18030"))})
    assert response.status_code == 201
    assert client.get(f"/api/datasets/{response.json()['id']}").json()["rows"][0]["编号"] == "00123"


def test_templates_and_invalid_execution(client):
    rules = [{"id": "1", "type": "trim", "enabled": False}]
    identity = client.get("/api/datasets").json()[0]["id"]
    assert client.post(f"/api/datasets/{identity}/preview", json={"rules": rules}).status_code == 400
    response = client.post("/api/templates", json={"name": "基础整理", "description": "用于台账", "rules": rules})
    assert response.status_code == 201
    assert client.get("/api/templates").json()[0]["name"] == "基础整理"
    assert client.delete(f"/api/templates/{response.json()['id']}").status_code == 204
    assert client.get("/api/templates").json() == []
    assert client.get("/api/datasets/nonexistent").status_code == 404
    assert client.get("/api/tasks/nonexistent/export").status_code == 404
    assert client.get(f"/api/datasets/{identity}?page=0").status_code == 400


def test_export_formula_is_neutralized(client):
    identity = upload_csv(client, '姓名,备注\n测试,=1+1\n金额,-12.50\n').json()["id"]
    task_id = client.post(f"/api/datasets/{identity}/tasks", json={"rules": [{"id": "1", "type": "trim"}]}).json()["id"]
    assert "'=1+1" in client.get(f"/api/tasks/{task_id}/export").text
    response = client.get(f"/api/tasks/{task_id}/export?format=xlsx")
    cell = openpyxl.load_workbook(io.BytesIO(response.content)).active["B2"]
    assert cell.data_type == "s"
    assert cell.value == "=1+1"
    assert "金额,-12.50" in client.get(f"/api/tasks/{task_id}/export").text


def test_preview_is_limited_but_saved_results_are_complete(client):
    text = "编号,城市\n" + "\n".join(f"{i:04d}, 济南 " for i in range(240))
    identity = upload_csv(client, text).json()["id"]
    body = {"rules": [{"id": "trim", "type": "trim"}]}
    preview = client.post(f"/api/datasets/{identity}/preview", json=body).json()
    assert len(preview["rows"]) == 200
    assert preview["after"]["rows"] == 240
    task_id = client.post(f"/api/datasets/{identity}/tasks", json=body).json()["id"]
    result = client.get(f"/api/tasks/{task_id}/result?page=5&page_size=50").json()
    assert len(result["items"]) == 40
    assert result["total"] == 240
    assert result["items"][-1]["_row_id"] == 240


def test_restart_marks_interrupted_tasks_and_does_not_reseed(client):
    store.query("INSERT INTO tasks VALUES (?,?,?,?,?,?,?,?,?,?)", ("interrupted", "中断任务", "missing", "文件.csv", "running", store.now(), None, "[]", None, None))
    count_before = len(client.get("/api/datasets").json())
    store.initialize()
    assert len(client.get("/api/datasets").json()) == count_before
    assert client.get("/api/tasks/interrupted").json()["status"] == "failed"
